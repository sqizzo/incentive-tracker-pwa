"""Buat supabase/seed_pricelist.sql dari Excel review price list penjualan.

Pemakaian (dari folder incentive-tracker-pwa/):
    python supabase/gen_pricelist_seed.py "../Price List Penjualan - Review.xlsx"

Sheet yang dibaca: "Harga Jual", "Pemetaan Item", "Pemetaan Unit" (format hasil review).
Seed bersifat idempoten: bisa dijalankan ulang untuk memperbarui harga.
File hasilnya berisi data bisnis dan di-ignore git; simpan di tempat privat.
"""

import os
import sys
import warnings
from pathlib import Path

from openpyxl import load_workbook

warnings.filterwarnings("ignore")

src = Path(sys.argv[1] if len(sys.argv) > 1 else "../Price List Penjualan - Review.xlsx")
out = Path(__file__).with_name("seed_pricelist.sql")
wb = load_workbook(src, data_only=True)

# Info price list: bisa di-override lewat environment variable
PL_NAME = os.environ.get("PL_NAME", "Price List Penjualan")
PL_EFFECTIVE = os.environ.get("PL_EFFECTIVE") or None  # format YYYY-MM-DD

NONE_MEMO = "(belum ada di memo)"
MIXED = "(campuran: dipilih saat input)"


def q(v):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return "null"
    if isinstance(v, (int, float)):
        return str(int(v))
    return "'" + str(v).strip().replace("'", "''") + "'"


def rows(ws, first=5):
    for r in range(first, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(v not in (None, "") for v in vals[:3]):
            yield vals


items = [(v[1], v[2], v[3]) for v in rows(wb["Pemetaan Item"])]  # kelompok, item, kategori
units = []
for v in rows(wb["Pemetaan Unit"]):
    key, kel, dev, note = v[0], v[1], v[2], v[3]
    if dev == NONE_MEMO:
        mode, dev = "none", None
    elif dev in (MIXED, None, ""):
        mode, dev = "choose", None
    else:
        mode = "fixed"
    units.append((key, kel, mode, dev, note))
prices = []
for v in rows(wb["Harga Jual"]):
    no, _hal, kel, unit, item, var, gar, harga, _status, _note, key = v[:11]
    # kolom Catatan berisi komentar review, tidak dibawa ke database
    prices.append((unit, kel, item, key, var, gar, harga, int(no), None))

sql = [
    f"-- Seed price list PENJUALAN. Digenerate oleh gen_pricelist_seed.py dari: {src.name}",
    f"-- {len(items)} item, {len(units)} unit, {len(prices)} harga. Jalankan setelah seed.sql (tarif insentif).",
    "begin;",
    "",
    "insert into public.pl_items (kelompok, name, category_id, sort_order)",
    "select v.kelompok, v.name, c.id, v.sort_order from (values",
    ",\n".join(f"  ({q(k)}, {q(i)}, {q(cat)}, {n + 1})" for n, (k, i, cat) in enumerate(items)),
    ") as v(kelompok, name, category_name, sort_order)",
    "left join public.service_categories c on c.name = v.category_name",
    "on conflict (kelompok, name) do update set category_id = excluded.category_id, sort_order = excluded.sort_order;",
    "",
    "insert into public.pl_units (unit_key, kelompok, incentive_mode, incentive_device_id, note)",
    "select v.unit_key, v.kelompok, v.mode, d.id, v.note from (values",
    ",\n".join(f"  ({q(k)}, {q(kel)}, {q(m)}, {q(d)}, {q(n)})" for k, kel, m, d, n in units),
    ") as v(unit_key, kelompok, mode, device_name, note)",
    "left join public.devices d on d.name = v.device_name",
    "on conflict (unit_key) do update set kelompok = excluded.kelompok, incentive_mode = excluded.incentive_mode,",
    "  incentive_device_id = excluded.incentive_device_id, note = excluded.note;",
    "",
    "-- harga: hapus lalu isi ulang agar sinkron dengan Excel (entri lama menyimpan snapshot, aman)",
    "delete from public.pl_prices;",
    "insert into public.pl_prices (unit_name, item_id, unit_id, variant, warranty, price, sort_order, note)",
    "select v.unit_name, i.id, u.id, v.variant, v.warranty, v.price, v.sort_order, v.note from (values",
    ",\n".join(
        f"  ({q(u)}, {q(kel)}, {q(it)}, {q(k)}, {q(var)}, {q(g)}, {q(h)}::int, {so}, {q(n)})"
        for u, kel, it, k, var, g, h, so, n in prices
    ),
    ") as v(unit_name, kelompok, item_name, unit_key, variant, warranty, price, sort_order, note)",
    "join public.pl_items i on i.kelompok = v.kelompok and i.name = v.item_name",
    "join public.pl_units u on u.unit_key = v.unit_key;",
    "",
    f"update public.rate_meta set sales_pl_name = {q(PL_NAME)}, sales_effective = {q(PL_EFFECTIVE)} where id = 1;",
    "",
    "commit;",
    "",
]
out.write_text("\n".join(sql), encoding="utf-8")
print(f"{out}: {len(items)} item, {len(units)} unit, {len(prices)} harga")
